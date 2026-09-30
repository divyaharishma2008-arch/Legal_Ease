import re
from io import BytesIO
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF


def sanitize_text(text: str) -> str:
    """Make text safe for the default PDF font while preserving normal content."""
    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2022": "-",
        "\u00a0": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"[^\x09\x0A\x0D\x20-\x7E]", "", text)
    return text.strip()


def format_txt(text: str) -> bytes:
    return text.encode("utf-8")


def _add_logo(doc: Document, logo_path: Optional[str]):
    if logo_path:
        try:
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(logo_path, width=Inches(1.25))
        except Exception:
            pass


def format_docx(
    text: str,
    doc_type: str,
    logo_path: Optional[str] = None,
) -> bytes:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    _add_logo(doc, logo_path)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            doc.add_paragraph()
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.08

        run = p.add_run(line)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

        if (
            len(line) < 90
            and (
                line.isupper()
                or re.match(r"^\d+[\.\)]\s+", line)
                or line.lower().startswith(("important notice", "parties", "signatures"))
            )
        ):
            run.bold = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("LegalEase - AI-generated draft")
    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(9)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


class LegalEasePDF(FPDF):
    def __init__(self, doc_type: str, logo_path: Optional[str] = None):
        super().__init__()
        self.doc_type = doc_type
        self.logo_path = logo_path

    def header(self):
        if self.logo_path:
            try:
                self.image(self.logo_path, x=95, y=8, w=20)
                self.ln(18)
            except Exception:
                pass
        self.set_font("Helvetica", "B", 12)
        self.cell(0, 8, self.doc_type.upper(), align="C")
        self.ln(10)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "", 8)
        self.cell(0, 10, "LegalEase - AI-generated draft", align="C")


def format_pdf(
    text: str,
    doc_type: str,
    logo_path: Optional[str] = None,
) -> bytes:
    pdf = LegalEasePDF(doc_type, logo_path)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)

    safe_text = sanitize_text(text)

    for raw_line in safe_text.splitlines():
        line = raw_line.strip()
        if not line:
            pdf.ln(4)
            continue

        is_heading = (
            line.isupper()
            or re.match(r"^\d+[\.\)]\s+", line)
            or line.lower().startswith(("important notice", "parties", "signatures"))
        )

        if is_heading:
            pdf.set_font("Helvetica", "B", 11)
        else:
            pdf.set_font("Helvetica", "", 11)

        pdf.multi_cell(0, 6, line)
        pdf.ln(1)

    return bytes(pdf.output())


def format_html_preview(text: str) -> str:
    """Create safe HTML for the Streamlit preview."""
    import html

    escaped = html.escape(text)
    paragraphs = []
    for block in escaped.split("\n\n"):
        block = block.replace("\n", "<br>")
        paragraphs.append(f"<p>{block}</p>")

    return f"""
    <div style="
        background:#111827;
        color:#f3f4f6;
        padding:24px;
        border-radius:14px;
        max-height:620px;
        overflow-y:auto;
        line-height:1.65;
        font-family:Georgia, 'Times New Roman', serif;
        border:1px solid #374151;">
        {''.join(paragraphs)}
    </div>
    """
